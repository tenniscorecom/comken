"""列記号（A, B, AA…）と列番号の相互変換のテスト。"""

import pytest

from comken.core.columns import col_to_num, column_number, num_to_col
from comken.exceptions import InvalidColumnError


class TestNumToCol:
    @pytest.mark.parametrize(
        ("number", "letter"),
        [
            (1, "A"),
            (2, "B"),
            (26, "Z"),
            (27, "AA"),
            (52, "AZ"),
            (53, "BA"),
            (702, "ZZ"),
            (703, "AAA"),
            (16384, "XFD"),
        ],
    )
    def test_known_values(self, number: int, letter: str) -> None:
        assert num_to_col(number) == letter

    def test_round_trip_with_col_to_num(self) -> None:
        """番号 → 列記号 → 番号で元に戻る（Excel の最大列 16384 を超える範囲まで）。"""
        for number in range(1, 20000):
            assert col_to_num(num_to_col(number)) == number

    def test_matches_openpyxl(self) -> None:
        """openpyxl の ``get_column_letter`` と同じ結果になる（comken がそこから置き換えたため）。"""
        from openpyxl.utils import get_column_letter

        for number in range(1, 18279):
            assert num_to_col(number) == get_column_letter(number)

    @pytest.mark.parametrize("bad", [0, -1, 1.5, "A", None, True])
    def test_invalid_number_raises(self, bad: object) -> None:
        with pytest.raises(InvalidColumnError):
            num_to_col(bad)  # type: ignore[arg-type]


class TestExistingConversions:
    def test_col_to_num_is_case_insensitive(self) -> None:
        assert col_to_num("aa") == 27

    def test_column_number_accepts_both(self) -> None:
        assert column_number("B") == 2
        assert column_number(2) == 2
